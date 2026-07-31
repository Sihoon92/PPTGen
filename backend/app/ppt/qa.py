"""Content QA on the rendered .pptx.

Re-opens the produced file with python-pptx (offline, reliable) and checks for the
failure modes a small model tends to produce: leftover placeholder tokens, empty
decks, and slide-count drift. Returns structured Issues; never raises.
"""

from __future__ import annotations

from app.ppt.dsl import Issue

PLACEHOLDER_TOKENS = ("{{", "}}", "lorem", "ipsum", "todo", "xxxx", "placeholder")


def _slide_texts(path: str) -> list[list[str]]:
    from pptx import Presentation

    prs = Presentation(path)
    out: list[list[str]] = []
    for slide in prs.slides:
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
        out.append(texts)
    return out


def run_content_qa(path: str | None, expected_slide_count: int) -> list[Issue]:
    if not path:
        return [Issue(stage="qa", code="no_output", message="no rendered file to inspect")]

    try:
        per_slide = _slide_texts(path)
    except Exception as exc:  # noqa: BLE001 - corrupt file etc.
        return [Issue(stage="qa", code="unreadable", message=f"cannot open pptx: {exc}")]

    issues: list[Issue] = []
    all_text = " ".join(t for slide in per_slide for t in slide).lower()

    if not all_text.strip():
        issues.append(Issue(stage="qa", code="empty_render", message="rendered deck has no text"))

    for token in PLACEHOLDER_TOKENS:
        if token in all_text:
            issues.append(
                Issue(stage="qa", code="placeholder",
                      message=f"placeholder token '{token}' left in output")
            )
            break

    if expected_slide_count and len(per_slide) != expected_slide_count:
        issues.append(
            Issue(stage="qa", severity="warning", code="slide_count",
                  message=f"rendered {len(per_slide)} slides, expected {expected_slide_count}")
        )

    return issues
