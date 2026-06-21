"""End-to-end smoke test: full PPT subgraph -> real Node renderer -> .pptx.

Gated on Node being available and the renderer's node_modules being installed
(``npm install`` in app/ppt/node_renderer). Skipped otherwise so CI without Node
stays green.
"""

import shutil

import pytest
from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
from langchain_core.messages import HumanMessage

from app.config import get_settings
from app.graph.nodes.ppt import build_ppt_subgraph
from app.ppt.renderer import SCRIPT_PATH

_NODE_MODULES = SCRIPT_PATH.parent / "node_modules" / "pptxgenjs"
pytestmark = pytest.mark.skipif(
    shutil.which("node") is None or not _NODE_MODULES.exists(),
    reason="node or pptxgenjs not installed (run `npm install` in app/ppt/node_renderer)",
)

DECK_SPEC_JSON = '{"title":"도입 효과","audience":"임원","goal":"설득","tone":"executive","narrative":["a","b"]}'
DECK_DSL_JSON = """[
  {"slide_id":"s1","role":"comparison","intent":"compare","title":"전략 비교",
   "layout":{"id":"cols","type":"columns","props":{"cols":2},"children":[
     {"id":"fast","type":"card","content":{"heading":"빠른 출시","body":{"type":"bullets","items":["MVP","저비용"]}}},
     {"id":"hybrid","type":"card","content":{"heading":"하이브리드","body":{"type":"bullets","items":["균형","추천"]}}}
   ]}},
  {"slide_id":"s2","role":"data_story","intent":"kpi","title":"핵심 지표",
   "layout":{"id":"kpis","type":"kpi_row","children":[
     {"id":"rt","type":"metric","content":{"label":"응답 시간","value":"-42%"}},
     {"id":"nps","type":"metric","content":{"label":"NPS","value":"+9"}}
   ]}}
]"""


@pytest.mark.asyncio
async def test_full_pipeline_produces_valid_pptx(tmp_path, monkeypatch):
    monkeypatch.setenv("ARTIFACTS_DIR", str(tmp_path / "artifacts"))
    get_settings.cache_clear()
    try:
        model = GenericFakeChatModel(messages=iter([DECK_SPEC_JSON, DECK_DSL_JSON]))
        graph = build_ppt_subgraph(model)
        out = await graph.ainvoke(
            {"messages": [HumanMessage("AI 도입 효과를 임원에게 설득하는 발표")], "session_id": "smoke"}
        )

        assert out.get("output_path"), f"render failed: {out.get('issues')}"
        from pptx import Presentation

        prs = Presentation(out["output_path"])
        assert len(prs.slides._sldIdLst) == 2
        all_text = " ".join(
            sh.text_frame.text
            for s in prs.slides
            for sh in s.shapes
            if sh.has_text_frame
        )
        assert "하이브리드" in all_text
        assert "-42%" in all_text
    finally:
        get_settings.cache_clear()
